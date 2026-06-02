from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from modules.google_sheets_sync import (
    RAW_SHEETS,
    build_overview_values,
    build_cockpit_values,
    rows_to_sheet_values,
)


HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
TITLE_FILL = PatternFill("solid", fgColor="1C5C5C")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
WHITE_FONT = Font(color="FFFFFF", bold=True)
HEADER_FONT = Font(bold=True)
BOLD_FONT = Font(bold=True)
THIN_SIDE = Side(style="thin", color="D0D7DE")
BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)


def safe(value):
    if value is None:
        return ""
    return str(value).strip()


def dedupe_preserve_order(values):
    seen = set()
    output = []
    for value in values or []:
        item = safe(value)
        if not item or item in seen:
            continue
        seen.add(item)
        output.append(item)
    return output


def chunked(values, size):
    for i in range(0, len(values), size):
        yield values[i:i + size]


def fetch_all_rows_paginated(supabase, table_name, chunk_size=1000):
    all_rows = []
    start = 0

    while True:
        end = start + chunk_size - 1
        response = (
            supabase.table(table_name)
            .select("*")
            .range(start, end)
            .execute()
        )
        rows = response.data or []

        if not rows:
            break

        all_rows.extend(rows)

        if len(rows) < chunk_size:
            break

        start += chunk_size

    return all_rows


def fetch_rows_for_ids(supabase, table_name, column_name, ids, chunk_size=100):
    ids = dedupe_preserve_order(ids)
    if not ids:
        return []

    collected = []
    seen_keys = set()

    try:
        for chunk in chunked(ids, chunk_size):
            response = (
                supabase.table(table_name)
                .select("*")
                .in_(column_name, chunk)
                .execute()
            )
            rows = response.data or []

            for row in rows:
                row_key = row.get("id") or tuple(sorted(row.items()))
                if row_key in seen_keys:
                    continue
                seen_keys.add(row_key)
                collected.append(row)

        if collected:
            return collected
    except Exception:
        pass

    wanted = set(ids)
    all_rows = fetch_all_rows_paginated(supabase, table_name)

    for row in all_rows:
        value = safe(row.get(column_name))
        if value not in wanted:
            continue

        row_key = row.get("id") or tuple(sorted(row.items()))
        if row_key in seen_keys:
            continue
        seen_keys.add(row_key)
        collected.append(row)

    return collected


def sort_rows(rows, keys):
    def sort_value(value):
        return safe(value).lower()

    return sorted(
        rows,
        key=lambda row: tuple(sort_value(row.get(k, "")) for k in keys),
    )


def write_values_to_sheet(ws, values, sheet_name):
    values = values or [["No data"]]

    for row in values:
        ws.append(list(row))

    if sheet_name == "Cockpit":
        style_cockpit_sheet(ws)
    else:
        style_normal_sheet(ws)

    auto_fit_columns(ws, max_width=45)


def style_normal_sheet(ws):
    if ws.max_row >= 1:
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    if ws.title == "Overview":
        for cell in ws["N"][1:]:
            cell.number_format = "0.0%"


def style_cockpit_sheet(ws):
    for cell in ws[1]:
        cell.fill = TITLE_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    for cell in ws[5]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    for row in range(6, ws.max_row + 1):
        ws[f"A{row}"].font = BOLD_FONT
        ws[f"A{row}"].alignment = Alignment(vertical="top", wrap_text=True)
        ws[f"B{row}"].alignment = Alignment(vertical="top", wrap_text=True)

    ws["B3"].fill = INPUT_FILL
    ws["B3"].border = BORDER
    ws.freeze_panes = "A6"


def auto_fit_columns(ws, max_width=45):
    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        max_len = 0

        for cell in col_cells:
            value = cell.value
            if value is None:
                continue

            if isinstance(value, str) and value.startswith("="):
                length = min(len(value), 20)
            else:
                length = len(str(value))

            if length > max_len:
                max_len = length

        ws.column_dimensions[col_letter].width = min(max_len + 2, max_width)


def build_filtered_workbook_bytes(supabase, register_ids, log_callback=None):
    register_ids = dedupe_preserve_order(register_ids)

    if not register_ids:
        raise ValueError("No register IDs were selected for export.")

    if log_callback:
        log_callback(f"Fetching rows for {len(register_ids)} companies...")

    companies = fetch_rows_for_ids(supabase, "companies", "register_id", register_ids)
    shareholders = fetch_rows_for_ids(supabase, "shareholders", "company_register_id", register_ids)
    company_news = fetch_rows_for_ids(supabase, "company_news", "company_register_id", register_ids)
    company_models = fetch_rows_for_ids(supabase, "company_models", "company_register_id", register_ids)
    fit_scores = fetch_rows_for_ids(supabase, "company_fit_scores", "company_register_id", register_ids)
    processing_logs = fetch_rows_for_ids(supabase, "processing_logs", "company_register_id", register_ids)

    companies = sort_rows(companies, ["register_id", "name"])
    shareholders = sort_rows(shareholders, ["company_register_id", "shareholder_name", "retrieved_at"])
    company_news = sort_rows(company_news, ["company_register_id", "date", "title", "retrieved_at"])
    company_models = sort_rows(company_models, ["company_register_id", "model_provider", "model_name", "updated_at", "created_at"])
    fit_scores = sort_rows(fit_scores, ["company_register_id", "model_provider", "model_name", "updated_at", "created_at"])
    processing_logs = sort_rows(processing_logs, ["company_register_id", "created_at", "module"])

    if log_callback:
        log_callback("Building workbook sheets...")

    overview_values = build_overview_values(
        companies=companies,
        shareholders=shareholders,
        company_news=company_news,
        company_models=company_models,
        fit_scores=fit_scores,
    )

    cockpit_values = build_cockpit_values()

    raw_data_map = {
        "North Data Exports": companies,
        "Shareholders": shareholders,
        "News": company_news,
        "Company Models": company_models,
        "Fit Scores": fit_scores,
        "Processing Logs": processing_logs,
    }

    table_counts = {
        "Overview": max(len(overview_values) - 1, 0),
        "North Data Exports": len(companies),
        "Shareholders": len(shareholders),
        "News": len(company_news),
        "Company Models": len(company_models),
        "Fit Scores": len(fit_scores),
        "Processing Logs": len(processing_logs),
    }

    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    sheet_order = ["Overview", "Cockpit"] + [cfg["sheet_name"] for cfg in RAW_SHEETS]

    for sheet_name in sheet_order:
        ws = wb.create_sheet(title=sheet_name)

        if sheet_name == "Overview":
            write_values_to_sheet(ws, overview_values, sheet_name)
        elif sheet_name == "Cockpit":
            write_values_to_sheet(ws, cockpit_values, sheet_name)
        else:
            config = next((cfg for cfg in RAW_SHEETS if cfg["sheet_name"] == sheet_name), None)
            rows = raw_data_map.get(sheet_name, [])
            values = rows_to_sheet_values(
                rows,
                exclude_columns=(config or {}).get("exclude_columns", []),
            )
            write_values_to_sheet(ws, values, sheet_name)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return {
        "workbook_bytes": buffer.getvalue(),
        "table_counts": table_counts,
        "selected_register_ids": register_ids,
        "company_rows": len(companies),
    }
